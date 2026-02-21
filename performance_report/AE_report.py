import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.borders import Border, Side
from openpyxl.utils import get_column_letter
from PyQt5.QtWidgets import QApplication, QFileDialog
import os
from performance_report.base_report import BaseBranchExcelProcessor


class AEBaseExcelProcessor(BaseBranchExcelProcessor):
    def __init__(self):
        super().__init__()

    def process_sheets(self):
        for sheet_name in self._input_dictionary.keys():
            input_dict = self._input_dictionary[sheet_name]
            output_sheet = self._output_workbook.create_sheet(title=sheet_name)
            output_sheet['A2'] = 'Employee Name'
            output_sheet['B2'] = input_dict['Employee Name']

            output_sheet['A3'] = 'Current Title'
            output_sheet['B3'] = input_dict['Current Title']

            output_sheet['A4'] = 'Date of Hire'
            output_sheet['B4'] = input_dict['Date of Hire']

            output_sheet['F2'] = 'Reviewer Name'
            output_sheet['G2'] = input_dict['Reviewer Name']

            output_sheet['F3'] = 'Review Date (Date of filling the form)'
            output_sheet['G3'] = input_dict['Review Date (Date of filling the form)']

            output_sheet['F4'] = 'Date of Last Review'
            output_sheet['G4'] = input_dict['Date of Last Review']

            output_sheet.merge_cells(start_row=2, start_column=3, end_row=4, end_column=5)

            output_sheet['A6'] = 'Parameters'
            output_sheet['B6'] = 'KPI'
            output_sheet['C6'] = 'Count'
            output_sheet['D6'] = 'Self Rating'
            output_sheet['E6'] = 'Reviewer Rating'
            output_sheet['F6'] = 'Weightage(%)'
            output_sheet['G6'] = 'Comments(Provide  examples for each category if possible)'
            output_sheet['A7'] = ('Technical Ability/Knowledge (To be assessed based on the following):\n1.  Increase in '
                         'technical skills; technical acumen progress rate\n2. Ability to solve technical issues with '
                         'minimal guidance\n3. Ability to learn and apply\n4. Up to date knowlege of interconnection '
                         'requirements/ regulations as related to buisness')
            output_sheet['A11'] = 'Quality/Accuracy of work/Timely Execution of Projects(To be assessed based on the following):\n1. Has an ongoing focus on accuracy/finds and corrects errors\n2. Minimal Rework for self\n3. Meeting projects due date (efficiency) '
            output_sheet['A15'] = 'Team Work(To be assessed based on the following):\n1. Willing to work with colleagues/team members diligently and efficiently  towards completing a project\n2. Clear communication with team on assigned tasks and deliverables'
            output_sheet['A16'] = '*Total hrs.\n1. Consulting (Billable and Non billable):\n2. Internal rework:\n3. Internal Admin/ Tech Disc/Open hr/Proposals:\n4. Internal training and development:\n5. Internal Research:\n(*Employee please leave blank. To be completed by Reviewer)'
            output_sheet['A17'] = 'Overall Summary (Explain in short how have you conributed towards the growth of the company)'
            output_sheet['A18'] = 'Goals for Next Year( Write in short atleast three goals, including any self development goals for next year)'
            output_sheet['A19'] = '*Overall Performance Rating ___________ and Feedback from Reviewer (Write in short performance feedback to the employee)\n(*Employee please leave blank. To be completed by Reviewer)'
            output_sheet['F7'] = '20'
            output_sheet['F11'] = '30'
            output_sheet['F15'] = '25'
            output_sheet['F16'] = '25'

            output_sheet['B7'] = 'How many Projects are done ? (Numerical)'
            output_sheet['B8'] = 'Type of projects (Please list all types of projects worked)'
            output_sheet['B9'] = 'Power System Tools (Please specify all power system tools used in the comments column)\n'
            output_sheet['B10'] = 'No of scripts developed (if any) (Please specify when script was used)'

            output_sheet['C7'] = input_dict['How many Projects are done ? (Numerical)']
            output_sheet['C8'] = input_dict['Type of projects (Please list all types of projects worked)']
            output_sheet['C9'] = input_dict['Power System Tools (Please specify all power system tools used in the comments column)\n']
            output_sheet['C10'] = input_dict['No of scripts developed (if any) (Please specify when script was used)']

            output_sheet['D7'] = input_dict['Self']
            output_sheet['E7'] = input_dict['Reviewer']
            output_sheet['G7'] = input_dict['Comments']

            output_sheet['B11'] = 'Number of projects done with zero deficiency (Numerical)'
            output_sheet['B12'] = 'Maintained timelines defined by PE/Client'
            output_sheet['B13'] = 'No of revisions required by Supervisor per project(average #) (Numerical)'

            output_sheet['C11'] = input_dict['Number of projects done with zero deficiency (Numerical)']
            output_sheet['C12'] = input_dict['Maintained timelines defined by PE/Client']
            output_sheet['C13'] = input_dict['No of revisions required by Supervisor per project(average #) (Numerical)']

            output_sheet['D11'] = input_dict['Self2']
            output_sheet['E11'] = input_dict['Reviewer2']
            output_sheet['G11'] = input_dict['Comments2']

            output_sheet['D15'] = input_dict['Self3']
            output_sheet['E15'] = input_dict['Reviewer3']
            output_sheet['G15'] = input_dict['Comments3']

            output_sheet['B17'] = input_dict['Explain in short how have you contributed towards the growth of the company']
            output_sheet['B18'] = input_dict[' Write in short at least\xa0three goals, including any self development goals for next year']
            output_sheet['B19'] = input_dict['Write in short performance feedback to the employee)                                             (*Employee please leave blank. To be completed by Reviewer']

    def apply_styles(self):
        for output_sheet_name in self._output_workbook.sheetnames:
            output_sheet = self._output_workbook[output_sheet_name]
            output_sheet.merge_cells(start_row=7, start_column=1, end_row=10, end_column=1)
            output_sheet['A7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=1, end_row=14, end_column=1)
            output_sheet['A11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=4, end_row=10, end_column=4)
            output_sheet['D7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=4, end_row=14, end_column=4)
            output_sheet['D11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=5, end_row=10, end_column=5)
            output_sheet['E7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=5, end_row=14, end_column=5)
            output_sheet['E11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=7, end_row=10, end_column=7)
            output_sheet['G7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=7, end_row=14, end_column=7)
            output_sheet['G11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=6, end_row=10, end_column=6)
            output_sheet['F7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=6, end_row=14, end_column=6)
            output_sheet['F11'].alignment = Alignment(vertical='center')

            green_cells = ['A2', 'A3', 'A4', 'F2', 'F3', 'F4', 'A6', 'B6', 'C6', 'D6', 'E6', 'F6', 'G6']
            for cell in green_cells:
                output_sheet[cell].fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
                output_sheet[cell].font = Font(color="FFFFFF")

            for row in output_sheet.iter_rows():
                for cell in row:
                    cell.border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )

    def save_output(self):
        super().save_output('AE')


# Example of how to use the class
# processor = AEBaseExcelProcessor()
# if processor.process_excel():
#     print("Processing complete. Output file saved to Desktop.")
# else:
#     print("No file selected or processing failed.")
