import os

import openpyxl
from PyQt5.QtWidgets import QApplication, QFileDialog, QInputDialog, QMessageBox


class BaseBranchExcelProcessor:
    def __init__(self):
        self.file_path = None
        self._input_dictionary = None
        self._output_workbook = None
        self._output_workbook_path = None

    # def select_file(self):
    #     app = QApplication([])
    #     self.file_path, _ = QFileDialog.getOpenFileName(None, "Select Excel File", "", "Excel files (*.xlsx;*.xls)")
    #     app.quit()
    #     return bool(self.file_path)
    @property
    def input_dictionary(self):
        return self._input_dictionary

    @input_dictionary.setter
    def input_dictionary(self, value):
        self._input_dictionary = value

    def input_dictionary_present(self):
        if self._input_dictionary is not None:
            return True
        else:
            return False

    @property
    def output_workbook(self):
        return self._output_workbook

    @output_workbook.setter
    def output_workbook(self, value):
        self._output_workbook = value

    @property
    def output_workbook_path(self):
        return self._output_workbook_path

    @output_workbook_path.setter
    def output_workbook(self, value):
        self._output_workbook_path = value

    def load_workbooks(self):
        self._output_workbook = openpyxl.Workbook()

    def process_sheets(self):
        pass
    def apply_styles(self):
        pass

    def save_output(self, output_file_name):
        _output_file_path = self._output_workbook_path
        # print(_output_file_path)
        output_file = output_file_name + '_EPR.xlsx'
        new_file_path = os.path.join(_output_file_path, output_file)
        # print(new_file_path)
        # try:
        self._output_workbook.save(new_file_path)
        msg = QMessageBox()
        msg.setIcon(QMessageBox.Information)
        msg.setText("Result has been saved in the respective folder.")
        msg.setWindowTitle("Saved")
        msg.exec_()

        # QMessageBox.information(self, 'Operation Successful', f'Output saved successfully to {new_file_path}!')
        # except Exception as e:
        #     # QMessageBox.warning(self, 'Error', f'Error occurred while saving output: {e}')
        #     warning_box = QMessageBox()
        #     warning_box.setIcon(QMessageBox.Warning)
        #     warning_box.setWindowTitle("Warning")
        #     warning_box.setText("Error occurred while saving output")
        #     warning_box.exec_()

    def prepare_output_excel(self):
        if self.input_dictionary_present():
            self.load_workbooks()
            self.process_sheets()
            self.apply_styles()
            # if ok:
            self.save_output()
            # else:
            #     QMessageBox.information(self, 'Operation Unsuccessful',"Please make it sure that you input excel sheet name")
            return True
        return False
