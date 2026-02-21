import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.borders import Border, Side
from openpyxl.utils import get_column_letter
from PyQt5.QtWidgets import QApplication, QFileDialog
import os
from performance_report.base_report import BaseBranchExcelProcessor


class HRBaseExcelProcessor(BaseBranchExcelProcessor):
    def __init__(self):
        super().__init__()

    def process_sheets(self):
        for sheet_name in self._input_dictionary.keys():
            input_dict = self._input_dictionary[sheet_name]
            output_sheet = self._output_workbook.create_sheet(title=sheet_name)
            output_sheet['A2'] = 'Employee Name'
            output_sheet['B2'] = input_dict['Employee Name']

            output_sheet['A3'] = 'Current Title'
            output_sheet['B3'] = input_dict['Current Title']

            output_sheet['A4'] = 'Date of Hire'
            output_sheet['B4'] = input_dict['Date of Hire']

            output_sheet['E2'] = 'Reviewer Name'
            output_sheet['F2'] = input_dict['Reviewer Name']

            output_sheet['E3'] = 'Review Date (Date of filling the form)'
            output_sheet['F3'] = input_dict['Review Date (Date of filling the form)']

            output_sheet['E4'] = 'Date of Last Review'
            output_sheet['F4'] = input_dict['Date of Last Review']

            output_sheet.merge_cells(start_row=2, start_column=3, end_row=4, end_column=4)

            output_sheet['A6'] = 'Parameters'
            output_sheet['B6'] = 'KPI'
            output_sheet['C6'] = 'Self Rating'
            output_sheet['D6'] = 'Reviewer Rating'
            output_sheet['E6'] = 'Weightage(%)'
            output_sheet['F6'] = 'Comments(Provide  examples for each category if possible)'
            output_sheet[
                'A7'] = 'Ability/Knowledge (To be assessed based on the following):\n1. Timely Execution of tasks assigned\n2. Multi Task Management/ Time Management\n3. Ability to learn and apply\n4. Event Management & Travel Management'
            output_sheet['A16'] = 'Quality/Accuracy of work/Timely Execution of Task'
            output_sheet[
                'A21'] = 'Team Work(To be assessed based on the following):\n1. Willing to work with colleagues/team members diligently and efficiently  towards completing a project\n2. Clear communication with team on assigned tasks and deliverables'
            output_sheet[
                'A22'] = '*Total hrs.\n1) Internal Admin Meetings/Administrative\n2) Internal training and development:\n3) Hiring \n(*Employee please leave blank. To be completed by Reviewer)'
            output_sheet[
                'A23'] = 'Overall Summary (Explain in short how have you conributed towards the growth of the company)'
            output_sheet[
                'A24'] = 'Goals for Next Year( Write in short atleast three goals, including any self development goals for next year)'
            output_sheet[
                'A25'] = '*Overall Performance Rating ___________ and Feedback from Reviewer (Write in short performance feedback to the employee)\n(*Employee please leave blank. To be completed by Reviewer)'
            output_sheet['E7'] = '40'
            output_sheet['E16'] = '30'
            output_sheet['E21'] = '30'

            output_sheet['B7'] = 'Recruitment & Induction'
            output_sheet['B8'] = 'Back Ground Verification'
            output_sheet['B9'] = 'Employee Welfare '
            output_sheet['B10'] = 'Training & Development'
            output_sheet['B11'] = 'Employee Engagement activities'
            output_sheet['B12'] = 'Cost Control'
            output_sheet['B13'] = 'Business Development'
            output_sheet['B14'] = 'Quarterly & Annually Meet Planning'
            output_sheet['B15'] = 'Statutory Compliance'

            output_sheet['C7'] = input_dict['Self']
            output_sheet['C8'] = input_dict['Self2']
            output_sheet['C9'] = input_dict['Self3']
            output_sheet['C10'] = input_dict['Self4']
            output_sheet['C11'] = input_dict['Self5']
            output_sheet['C12'] = input_dict['Self6']
            output_sheet['C13'] = input_dict['Self7']
            output_sheet['C14'] = input_dict['Self8']
            output_sheet['C15'] = input_dict['Self9']

            output_sheet['D7'] = input_dict['Reviewer']
            output_sheet['D8'] = input_dict['Reviewer2']
            output_sheet['D9'] = input_dict['Reviewer3']
            output_sheet['D10'] = input_dict['Reviewer4']
            output_sheet['D11'] = input_dict['Reviewer5']
            output_sheet['D12'] = input_dict['Reviewer6']
            output_sheet['D13'] = input_dict['Reviewer7']
            output_sheet['D14'] = input_dict['Reviewer8']
            output_sheet['D15'] = input_dict['Reviewer9']

            output_sheet['F7'] = input_dict['Comments (Provide Examples for each Category if possible)']

            output_sheet['B16'] = 'Maintained timeline '
            output_sheet['B17'] = 'Employee Satisfaction'
            output_sheet['B18'] = 'Route Cause Analysis'
            output_sheet['B19'] = 'Control Billing as per the Usage'

            output_sheet['C16'] = input_dict['Self10']
            output_sheet['C17'] = input_dict['Self11']
            output_sheet['C18'] = input_dict['Self12']
            output_sheet['C19'] = input_dict['Self13']

            output_sheet['D16'] = input_dict['Reviewer10']
            output_sheet['D17'] = input_dict['Reviewer11']
            output_sheet['D18'] = input_dict['Reviewer12']
            output_sheet['D19'] = input_dict['Reviewer13']

            output_sheet['F16'] = input_dict['Comments (Provide Examples for each Category if possible)2']

            output_sheet['C21'] = input_dict['Self14']
            output_sheet['D21'] = input_dict['Reviewer14']
            output_sheet['F21'] = input_dict['Comments (Provide Examples)']

            output_sheet['B23'] = input_dict[('Explain in short how have you conributed towards the growth of the '
                                               'company')]
            output_sheet['B24'] = input_dict[('\xa0Write in short atleast three goals, including any self development '
                                               'goals for next year')]
            output_sheet['B25'] = input_dict[('Write in short performance feedback to the employee)\xa0(*Employee '
                                               'please leave blank. To be completed by Reviewer')]

    def apply_styles(self):
        for output_sheet_name in self._output_workbook.sheetnames:
            output_sheet = self._output_workbook[output_sheet_name]
            output_sheet.merge_cells(start_row=7, start_column=1, end_row=15, end_column=1)
            output_sheet['A7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=16, start_column=1, end_row=20, end_column=1)
            output_sheet['A16'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=4, end_row=10, end_column=4)
            output_sheet['D7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=4, end_row=14, end_column=4)
            output_sheet['D11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=5, end_row=10, end_column=5)
            output_sheet['E7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=5, end_row=14, end_column=5)
            output_sheet['E11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=7, end_row=10, end_column=7)
            output_sheet['G7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=7, end_row=14, end_column=7)
            output_sheet['G11'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=7, start_column=6, end_row=10, end_column=6)
            output_sheet['F7'].alignment = Alignment(vertical='center')

            output_sheet.merge_cells(start_row=11, start_column=6, end_row=14, end_column=6)
            output_sheet['F11'].alignment = Alignment(vertical='center')

            green_cells = ['A2', 'A3', 'A4', 'F2', 'F3', 'F4', 'A6', 'B6', 'C6', 'D6', 'E6', 'F6']
            for cell in green_cells:
                output_sheet[cell].fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
                output_sheet[cell].font = Font(color="FFFFFF")

            for row in output_sheet.iter_rows():
                for cell in row:
                    cell.border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin')
                    )

    def save_output(self):
        super().save_output('HR')


# Example of how to use the class
# processor = HRBaseExcelProcessor()
# if processor.process_excel():
#     print("Processing complete. Output file saved to Desktop.")
# else:
#     print("No file selected or processing failed.")
